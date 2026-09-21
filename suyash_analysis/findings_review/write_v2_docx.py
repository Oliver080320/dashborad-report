from pathlib import Path
import sys,tempfile,json,re,shutil
sys.path.insert(0,str(Path(tempfile.gettempdir())/'sales-report-tools'))
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
P=Path(__file__).resolve().parent
C=json.loads((P/'report_app/src/content/report/findings-v2.json').read_text(encoding='utf-8'))
S=json.loads((P/'report_app/src/data.json').read_text(encoding='utf-8'))
def setup():
 d=Document();sec=d.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
 sec.top_margin=sec.bottom_margin=Inches(.58);sec.left_margin=sec.right_margin=Inches(.65)
 st=d.styles['Normal'];st.font.name='Calibri';st.font.size=Pt(10);st.paragraph_format.space_after=Pt(5);st.paragraph_format.line_spacing=1.02
 for name,size in [('Title',21),('Heading 1',17),('Heading 2',12),('Heading 3',11)]:
  st=d.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('1F6F5C');st.paragraph_format.space_before=Pt(7);st.paragraph_format.space_after=Pt(5)
 h=sec.header.paragraphs[0];h.text='SALES INSIGHTS  /  REVIEW DRAFT  /  21 SEPTEMBER 2026';h.style='Caption';h.runs[0].font.size=Pt(8)
 f=sec.footer.paragraphs[0];f.text='Suyash review before circulation  •  v2                                      '
 fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
 d.core_properties.title=C['title'];d.core_properties.subject='Revised findings with preserved source evidence';d.core_properties.author='Sales analysis'
 return d
def rich(p,text):
 text=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1 (\2)',text)
 for i,part in enumerate(re.split(r'\*\*(.*?)\*\*',text)):
  r=p.add_run(part.replace('`',''));r.bold=bool(i%2)
 return p
def md(d,text):
 lines=text.splitlines();i=0
 while i<len(lines):
  line=lines[i].strip();i+=1
  if not line:continue
  if line.startswith('|'):
   rr=[]
   while True:
    if not re.fullmatch(r'[| :\-]+',line):rr.append([v.strip() for v in line.strip('|').split('|')])
    if i>=len(lines) or not lines[i].strip().startswith('|'):break
    line=lines[i].strip();i+=1
   if rr:table(d,rr[0],rr[1:])
  elif line.startswith('#'):rich(d.add_paragraph(style='Heading '+str(min(len(line)-len(line.lstrip('#')),3))),line.lstrip('# '))
  else:rich(d.add_paragraph(style='List Bullet' if line.startswith('- ') else None),line.removeprefix('- '))
def table(d,cols,rows):
 t=d.add_table(rows=1,cols=len(cols));t.style='Light Shading Accent 1'
 for c,v in zip(t.rows[0].cells,cols):rich(c.paragraphs[0],str(v));c.paragraphs[0].runs[0].bold=True
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for rr in rows:
  for c,v in zip(t.add_row().cells,rr):rich(c.paragraphs[0],str(v))
 for row in t.rows:
  no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
  for cell in row.cells:
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EEF4F1');cell._tc.get_or_add_tcPr().append(shade)
   for para in cell.paragraphs:
    para.paragraph_format.space_after=Pt(4);para.paragraph_format.space_before=Pt(3)
    for r in para.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('233B33')
 return t
def fig(d,id,maxheight=2.3):
 path=P/'v2_charts'/f'{id}.png';w,h=Image.open(path).size;width=min(6.97,maxheight*w/h)
 para=d.add_paragraph();para.paragraph_format.space_after=Pt(3);para.alignment=1;para.add_run().add_picture(str(path),width=Inches(width))
def score(d):
 t=d.add_table(rows=2,cols=3)
 for cell,k in zip([c for r in t.rows for c in r.cells],C['kpis']):
  shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EEF4F1');cell._tc.get_or_add_tcPr().append(shade)
  p=cell.paragraphs[0];r=p.add_run(k['label']);r.bold=True;r.font.size=Pt(9)
  r=cell.add_paragraph().add_run(k['value']);r.bold=True;r.font.size=Pt(23);r.font.color.rgb=RGBColor.from_string('1F6F5C')
  cell.add_paragraph(k['detail']).runs[0].font.size=Pt(8)
 d.add_paragraph().paragraph_format.space_after=Pt(0)
d=setup();markdown=['# Sales insights report — v2',C['subtitle']]
for idx,page in enumerate(C['pages']):
 if idx:d.add_page_break()
 if not idx:d.add_heading('Sales insights report',0)
 p=d.add_paragraph(page['section'].upper());p.runs[0].font.size=Pt(9);p.runs[0].bold=True
 d.add_heading(page['title'],1);markdown+=['## '+page['section'],page['title']]
 for b in page['blocks']:
  if b['type']=='kpis':score(d)
  elif b['type']=='text':md(d,b['text']);markdown.append(b['text'])
  elif b['type']=='table':table(d,b['columns'],b['rows']);markdown+=['| '+' | '.join(b['columns'])+' |','| '+' | '.join(['---']*len(b['columns']))+' |']+['| '+' | '.join(r)+' |' for r in b['rows']]
  else:
   d.add_heading(b['title'],2);md(d,b['text'])
   if b.get('figure'):fig(d,b['figure'],2.5 if page['id'] in ['team','wins','trends'] else 3.1)
   md(d,'**Action & owner:** '+b['action']);md(d,'**Confidence: '+b['confidence']+'** — '+b['reason'])
   markdown+=['### '+b['title'],b['text'],'**Action & owner:** '+b['action'],'**Confidence: '+b['confidence']+'** — '+b['reason']]
# Evidence appendix: each issue is count + denominator + examples + full extract path.
d.add_page_break();d.add_heading(C['appendices'][0]['title'],1);md(d,C['appendices'][0]['text'])
markdown+=['## '+C['appendices'][0]['title'],C['appendices'][0]['text']]
for issue in C['issues']:
 d.add_heading(issue['display_id']+' — '+issue['title'],2)
 txt=f"**Affected:** {issue['count']:,} / {issue['denominator']:,} ({issue['percent']:.2f}%). **Severity:** {issue['severity']}.\n\n{issue['meaning']}\n\n**Recommendation:** {issue['action']}"
 md(d,txt);examples=issue.get('example',[])[:2]
 for j,ex in enumerate(examples):md(d,'**Example '+str(j+1)+':** '+'; '.join(f'{k}={v}' for k,v in ex.items()))
 md(d,'**Complete proof:** evidence/'+issue['original_id']+'.csv. Original reference: '+issue['original_id']+'.')
 markdown+=['### '+issue['display_id']+' — '+issue['title'],txt,'Proof: evidence/'+issue['original_id']+'.csv']
d.add_page_break();d.add_heading(C['appendices'][1]['title'],1);md(d,C['appendices'][1]['text'])
md(d,'**V2 comparisons:** chart-ready CSVs are in evidence_v2. January–July: 270 / 2,105 = 12.83%. July–August change: (362 − 439) / 439 = −17.54%. V2 external customer rankings exclude internal organisation #91445. Displayed percentages are rounded; underlying extracts preserve precision.')
inventory=S['queries']['source_inventory']['rows']
for r in inventory:
 md(d,'**'+str(r.get('file'))+'** — SHA256: '+str(r.get('sha256')))
markdown+=['## '+C['appendices'][1]['title'],C['appendices'][1]['text']]
d.add_page_break();d.add_heading(C['appendices'][2]['title'],1);md(d,C['appendices'][2]['text'])
for support in C['support']:
 d.add_heading(support['title'],2);md(d,support['text']);markdown+=['### '+support['title'],support['text']]
d.add_heading('Dashboard panel evidence',2)
for panel in C['panels']:
 md(d,'**'+panel['panel']+'** — '+panel['finding']+'\n\nProof: evidence/'+panel['evidence_query']+'.csv')
md(d,'The companion evidence directory contains full extracts and original screenshots. archive_v1 preserves the previous report unchanged. evidence_v2/issue_crosswalk.csv maps all 22 original issues to their revised location; the four implementation findings and internal-account scope note are in Suyash_Dashboard_Fixes_and_Review_Notes.docx. No source records were edited.')
d.save(P/'Sales_Insights_Report_v2.docx')
(P/'SALES_FINDINGS_REPORT.md').write_text('\n\n'.join(markdown),encoding='utf-8')
n=setup();md(n,(P/'Suyash_Dashboard_Fixes_and_Review_Notes.md').read_text(encoding='utf-8'));n.save(P/'Suyash_Dashboard_Fixes_and_Review_Notes.docx')
print('Wrote revised report and separate Suyash note.')
